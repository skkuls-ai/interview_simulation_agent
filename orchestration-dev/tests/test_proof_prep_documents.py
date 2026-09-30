"""서류 접수·텍스트 추출 테스트 (docs/03 F-001, docs/08 T-107·T-108·T-109·T-209·T-210)."""

from __future__ import annotations

import io
import json
from datetime import datetime, timezone
from pathlib import Path

import docx
import pytest
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

from backend.proof_prep.documents import DOC_FIELDS, DocumentError, clean_text, collect_documents, extract_text

DEMO = Path(__file__).parents[1] / "data" / "demo"
SAMPLE = json.loads((DEMO / "sample_inputs.json").read_text(encoding="utf-8"))
TEXTS = {f: SAMPLE[f"{f}_text"] for f in DOC_FIELDS}
LINE = "레시피 RAG 챗봇에서 전처리 규칙을 바꿔 검색 정확도 20% 개선"


def make_pdf(*lines: str) -> bytes:
    pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.setFont("HYSMyeongJo-Medium", 11)
    for i, line in enumerate(lines):
        c.drawString(50, 800 - 20 * i, line)
    c.save()
    return buf.getvalue()


def make_blank_pdf() -> bytes:  # 글자가 없는 PDF (이미지 스캔본과 같은 상황)
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.rect(50, 50, 200, 200)
    c.save()
    return buf.getvalue()


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def assert_error(exc: pytest.ExceptionInfo, code: str, field: str | None = None):
    e = exc.value
    assert e.code == code and e.field == field
    body = e.to_body()["error"]
    assert body["code"] == code and body["message"]
    assert body.get("field") == field


# ================================================================ 형식별 추출


def test_pdf_korean_text_is_extracted():
    text = extract_text("resume.pdf", make_pdf(LINE, "FastAPI와 React로 채팅 UI 구현 (4인 팀)"), field="resume")
    assert LINE in text and "4인 팀" in text


def test_docx_paragraphs_and_tables_are_extracted():
    data = make_docx([LINE], table=[["기간", "프로젝트"], ["2026.09", "ProofLetter 멀티 에이전트 (팀 리드)"]])
    text = extract_text("resume.docx", data, field="resume")
    assert LINE in text
    assert "2026.09 | ProofLetter 멀티 에이전트 (팀 리드)" in text  # 표는 행 단위로


def test_txt_utf8_and_windows_korean_encoding():
    assert extract_text("a.txt", LINE.encode("utf-8")) == LINE
    assert extract_text("b.txt", ("﻿" + LINE).encode("utf-8")) == LINE  # BOM
    assert extract_text("c.txt", LINE.encode("cp949")) == LINE              # Windows 메모장


def test_text_is_cleaned_but_lines_are_kept():
    raw = "담당 업무\r\n-   LangGraph 기반    멀티 에이전트 설계\r\n\r\n\r\n\r\n- FastAPI 서비스 API 개발"
    text = extract_text("jd.txt", raw.encode("utf-8"))
    assert text == "담당 업무\n- LangGraph 기반 멀티 에이전트 설계\n\n- FastAPI 서비스 API 개발"


# ================================================================ 추출 실패 (T-107, T-108)


@pytest.mark.parametrize("filename,data,reason", [
    ("scan.pdf", make_blank_pdf(), "스캔"),
    ("broken.pdf", b"%PDF-1.4 not really a pdf", "읽을 수 없습니다"),
    ("resume.hwp", b"HWP Document File", "HWP"),
    ("photo.png", b"\x89PNG....", "지원하지 않는 형식"),
    ("empty.txt", b"", "빈 파일"),
    ("short.txt", "이력서".encode(), "거의 추출하지 못했습니다"),
])
def test_unreadable_files_ask_for_text_input(filename, data, reason):
    with pytest.raises(DocumentError) as exc:
        extract_text(filename, data, field="resume")
    assert_error(exc, "TEXT_EXTRACTION_FAILED", "resume")
    assert reason in exc.value.message and "직접 입력" in exc.value.message


def test_too_large_file_is_rejected():
    with pytest.raises(DocumentError) as exc:
        extract_text("big.txt", b"a" * (10 * 1024 * 1024 + 1), field="cover_letter")
    assert_error(exc, "TEXT_EXTRACTION_FAILED", "cover_letter")


# ================================================================ 서류 4종 접수 (T-109, T-209, T-210)


def test_all_text_input_returns_state_fields():
    now = datetime(2026, 9, 30, 12, 0, tzinfo=timezone.utc)
    docs = collect_documents(texts=TEXTS, privacy_consent="true", now=now)
    fields = docs.as_state_fields()
    # 직접 입력한 글도 파일과 같은 규칙으로 정리합니다 (앞뒤 공백, 줄 끝 공백)
    assert fields["job_description_text"] == clean_text(TEXTS["job_description"])
    assert fields["consent_at"] == "2026-09-30T12:00:00+00:00"
    assert set(fields) == {"resume_text", "job_posting_text", "job_description_text", "cover_letter_text", "consent_at"}


def test_file_and_text_can_be_mixed():
    texts = {**TEXTS, "resume": None}
    files = {"resume": ("resume.pdf", make_pdf(LINE, "FastAPI와 React로 채팅 UI 구현 (4인 팀)"))}
    docs = collect_documents(files=files, texts=texts, privacy_consent=True)
    assert LINE in docs.resume_text
    assert docs.cover_letter_text == clean_text(TEXTS["cover_letter"])


def test_text_wins_when_both_file_and_text_are_given():
    files = {"resume": ("resume.txt", "파일에서 온 이력서 내용입니다. 서른 글자를 넘기기 위한 문장입니다.".encode())}
    docs = collect_documents(files=files, texts=TEXTS, privacy_consent=True)
    assert docs.resume_text == clean_text(TEXTS["resume"])


def test_consent_is_checked_first():
    with pytest.raises(DocumentError) as exc:
        collect_documents(texts={}, privacy_consent="false")  # 서류도 없지만 동의부터 봄
    assert_error(exc, "CONSENT_REQUIRED")


@pytest.mark.parametrize("missing", DOC_FIELDS)
def test_each_missing_document_is_named(missing):
    texts = {**TEXTS, missing: "   "}
    with pytest.raises(DocumentError) as exc:
        collect_documents(texts=texts, privacy_consent=True)
    assert_error(exc, "MISSING_REQUIRED_DOC", missing)


def test_extraction_failure_names_the_document():
    files = {"job_description": ("jd.pdf", make_blank_pdf())}
    texts = {**TEXTS, "job_description": None}
    with pytest.raises(DocumentError) as exc:
        collect_documents(files=files, texts=texts, privacy_consent=True)
    assert_error(exc, "TEXT_EXTRACTION_FAILED", "job_description")


def test_too_short_typed_text_is_rejected():
    with pytest.raises(DocumentError) as exc:
        collect_documents(texts={**TEXTS, "cover_letter": "열심히 하겠습니다."}, privacy_consent=True)
    assert_error(exc, "TEXT_EXTRACTION_FAILED", "cover_letter")
