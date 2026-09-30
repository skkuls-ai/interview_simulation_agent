"""업로드 문서 저장, 텍스트 추출, 삭제.

원칙: 문서 원본과 추출 텍스트는 세션 폴더에만 두고, 세션이 끝나면 폴더째 삭제합니다.
"""

from __future__ import annotations

import io
import shutil
from pathlib import Path
from typing import Literal

DocKind = Literal["jd", "resume", "cover_letter"]
DOC_LABELS = {"jd": "채용공고(JD)", "resume": "이력서", "cover_letter": "자기소개서"}
MAX_BYTES = 10 * 1024 * 1024
MAX_CHARS = 30_000
ALLOWED = {".pdf", ".docx", ".txt", ".md"}


class DocumentError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED:
        raise DocumentError(f"지원하지 않는 형식입니다 ({ext}). PDF, DOCX, TXT 파일을 올려 주세요.")
    if len(data) > MAX_BYTES:
        raise DocumentError("파일이 10MB를 넘습니다.")
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
        elif ext == ".docx":
            import docx

            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
        else:
            text = data.decode("utf-8-sig", errors="replace")
    except DocumentError:
        raise
    except Exception as e:  # 손상된 파일 등
        raise DocumentError(f"파일을 읽을 수 없습니다: {e}") from e

    text = "\n".join(line.rstrip() for line in text.splitlines() if line.strip())
    if len(text) < 30:
        raise DocumentError("텍스트를 거의 추출하지 못했습니다. 스캔한 이미지 PDF라면 텍스트가 있는 파일로 올려 주세요.")
    return text[:MAX_CHARS]


class DocumentStore:
    def __init__(self, root: Path):
        self.root = Path(root)
        self.root.mkdir(parents=True, exist_ok=True)

    def _dir(self, session_id: str) -> Path:
        return self.root / session_id

    def save(self, session_id: str, kind: DocKind, filename: str, data: bytes) -> str:
        text = extract_text(filename, data)
        d = self._dir(session_id)
        d.mkdir(parents=True, exist_ok=True)
        for old in d.glob(f"{kind}.*"):
            old.unlink()
        (d / f"{kind}{Path(filename).suffix.lower()}").write_bytes(data)
        (d / f"{kind}.extracted.txt").write_text(text, encoding="utf-8")
        return text

    def texts(self, session_id: str) -> dict[str, str]:
        d = self._dir(session_id)
        return {p.name.split(".")[0]: p.read_text(encoding="utf-8") for p in d.glob("*.extracted.txt")} if d.exists() else {}

    def kinds(self, session_id: str) -> list[str]:
        return sorted(self.texts(session_id))

    def delete_session(self, session_id: str) -> bool:
        d = self._dir(session_id)
        if d.exists():
            shutil.rmtree(d)
            return True
        return False

    def exists(self, session_id: str) -> bool:
        d = self._dir(session_id)
        return d.exists() and any(d.iterdir())
