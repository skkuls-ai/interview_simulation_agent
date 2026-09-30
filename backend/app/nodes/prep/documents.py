"""서류 접수·텍스트 추출 (담당 A, docs/03 F-001 · docs/04 §3.1).

C 의 POST /api/interviews 핸들러가 multipart 로 받은 값을 그대로 넘기면
서류 4종 텍스트와 동의 시각을 돌려주고, 문제가 있으면 DocumentError 를 냅니다.

    try:
        docs = collect_documents(files={"resume": (filename, data), ...},
                                 texts={"job_posting": "...", ...},
                                 privacy_consent=form.get("privacy_consent"))
    except DocumentError as e:
        return JSONResponse(status_code=422, content=e.to_body())

검사 순서: 동의 → 필수 서류 → 텍스트 추출 (docs/08 T-109, T-209, T-107)
"""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

# 서류 칸 이름. API 필드는 <이름>_file / <이름>_text (docs/04 §3.1)
DOC_FIELDS = ("resume", "job_posting", "job_description", "cover_letter")
DOC_LABELS = {"resume": "이력서", "job_posting": "채용공고", "job_description": "직무기술서", "cover_letter": "자기소개서"}

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_FILE_BYTES = 10 * 1024 * 1024  # 10MB (화면 2 안내 문구와 같게)
MAX_CHARS = 30_000                  # 너무 긴 서류는 앞부분만 (LLM 입력 한도)
MIN_CHARS = 30                      # (제안) 이보다 적게 추출되면 스캔 PDF 등으로 보고 텍스트 입력을 안내. docs 미정


class DocumentError(ValueError):
    """422 로 돌려줄 입력 오류. field 가 있으면 화면이 그 칸을 텍스트 입력으로 바꿉니다."""

    def __init__(self, code: str, message: str, field: str | None = None):
        super().__init__(message)
        self.code = code
        self.message = message
        self.field = field

    def to_body(self) -> dict:
        err = {"code": self.code, "message": self.message}
        if self.field:
            err["field"] = self.field
        return {"error": err}


@dataclass(frozen=True)
class CollectedDocuments:
    resume_text: str
    job_posting_text: str
    job_description_text: str
    cover_letter_text: str
    consent_at: str  # ISO-8601

    def as_state_fields(self) -> dict:
        """InterviewState 에 그대로 넣을 수 있는 dict (docs/04 State 필드명)."""
        return {
            "resume_text": self.resume_text,
            "job_posting_text": self.job_posting_text,
            "job_description_text": self.job_description_text,
            "cover_letter_text": self.cover_letter_text,
            "consent_at": self.consent_at,
        }


# ================================================================ 공개 함수


def collect_documents(
    files: dict[str, tuple[str, bytes] | None] | None = None,
    texts: dict[str, str | None] | None = None,
    privacy_consent: object = None,
    now: datetime | None = None,
) -> CollectedDocuments:
    """서류 4종을 받아 텍스트로 만듭니다.

    files: {"resume": (파일명, 내용 bytes), ...}   texts: {"resume": "직접 입력한 글", ...}
    한 칸에 파일과 텍스트가 둘 다 오면 텍스트를 씁니다 (제안: 사용자가 직접 고친 내용일 가능성이 높음).
    """
    files, texts = files or {}, texts or {}
    if not _is_true(privacy_consent):
        raise DocumentError("CONSENT_REQUIRED", "개인정보 수집·이용에 동의해야 분석을 시작할 수 있습니다.")

    for f in DOC_FIELDS:  # 먼저 빠진 칸이 있는지 전부 확인 (추출은 느리므로 나중에)
        if not (texts.get(f) or "").strip() and not files.get(f):
            raise DocumentError("MISSING_REQUIRED_DOC", f"{DOC_LABELS[f]}을(를) 입력해 주세요.", field=f)

    out = {}
    for f in DOC_FIELDS:
        typed = (texts.get(f) or "").strip()
        out[f] = clean_text(typed) if typed else extract_text(*files[f], field=f)
        if len(out[f]) < MIN_CHARS:
            raise DocumentError(
                "TEXT_EXTRACTION_FAILED",
                f"{DOC_LABELS[f]} 내용이 너무 짧습니다. 내용을 확인하거나 직접 입력해 주세요.", field=f,
            )

    consent_at = (now or datetime.now(timezone.utc)).isoformat(timespec="seconds")
    return CollectedDocuments(
        resume_text=out["resume"], job_posting_text=out["job_posting"],
        job_description_text=out["job_description"], cover_letter_text=out["cover_letter"], consent_at=consent_at,
    )


def extract_text(filename: str, data: bytes, field: str | None = None) -> str:
    """파일 하나에서 텍스트를 뽑습니다. 실패하면 TEXT_EXTRACTION_FAILED (field 포함)."""
    label = DOC_LABELS.get(field or "", "서류")
    ext = Path(filename or "").suffix.lower()

    def fail(reason: str) -> DocumentError:
        return DocumentError("TEXT_EXTRACTION_FAILED", f"{label}: {reason} 텍스트로 직접 입력해 주세요.", field=field)

    if ext == ".hwp" or ext == ".hwpx":
        raise fail("HWP 파일은 지원하지 않습니다. PDF·DOCX·TXT로 저장하거나")
    if ext not in ALLOWED_EXTENSIONS:
        raise fail(f"지원하지 않는 형식입니다({ext or '확장자 없음'}). PDF·DOCX·TXT 파일을 올리거나")
    if not data:
        raise fail("빈 파일입니다.")
    if len(data) > MAX_FILE_BYTES:
        raise fail("파일이 10MB를 넘습니다.")

    try:
        if ext == ".pdf":
            raw = _read_pdf(data)
        elif ext == ".docx":
            raw = _read_docx(data)
        else:
            raw = _read_txt(data)
    except DocumentError:
        raise
    except Exception:  # 손상된 파일, 암호 걸린 PDF 등
        raise fail("파일을 읽을 수 없습니다. 파일이 손상되었거나 암호가 걸려 있을 수 있습니다.") from None

    text = clean_text(raw)
    if len(text) < MIN_CHARS:
        raise fail("텍스트를 거의 추출하지 못했습니다. 이미지로 스캔한 PDF라면")
    return text


def clean_text(text: str) -> str:
    """줄 구조는 유지합니다 (분석이 글머리표·줄 단위로 읽기 때문). 인용 검증과 같은 NFC 정규화를 합니다."""
    t = unicodedata.normalize("NFC", text).replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t ]+", " ", line).strip() for line in t.split("\n")]
    t = re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
    return t[:MAX_CHARS]


# ================================================================ 형식별 읽기


def _read_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ValueError("encrypted")
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _read_docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:  # 이력서는 표로 된 경우가 많음
        for row in table.rows:
            cells = list(dict.fromkeys(c.text.strip() for c in row.cells))  # 병합 셀 중복 제거
            parts.append(" | ".join(c for c in cells if c))
    return "\n".join(parts)


def _read_txt(data: bytes) -> str:
    # Windows 메모장의 한글 TXT 는 cp949 인 경우가 많습니다.
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _is_true(value: object) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in ("true", "1", "yes", "on")
