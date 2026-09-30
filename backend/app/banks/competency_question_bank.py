"""Read the structured competency interview question workbook."""

from __future__ import annotations

import re
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

WORKBOOK_PATH = Path(__file__).with_name("역량기반_구조화_면접_질문_150선.xlsx")
EXPECTED_HEADERS = (
    "No", "대분류", "소분류", "소분류 설명", "질문유형", "질문", "추가 질문",
    "질문 의도", "Positive 체크포인트", "Negative 체크포인트",
)


def _lines(value: object) -> tuple[str, ...]:
    return tuple(line.strip() for line in str(value or "").splitlines() if line.strip())


@dataclass(frozen=True)
class CompetencyQuestion:
    number: int
    category: str
    competency: str
    competency_description: str
    question_type: str
    text: str
    follow_ups: tuple[str, ...]
    intents: tuple[str, ...]
    positive_checkpoints: tuple[str, ...]
    negative_checkpoints: tuple[str, ...]

    @property
    def question_bank_id(self) -> str:
        slug = re.sub(r"[^0-9A-Za-z가-힣]+", "-", self.competency).strip("-")
        return f"COMP-{slug}-{self.number:03d}"

    def selection_payload(self) -> dict[str, object]:
        return {
            "number": self.number,
            "category": self.category,
            "competency": self.competency,
            "competency_description": self.competency_description,
            "question_type": self.question_type,
            "question": self.text,
            "intents": list(self.intents),
        }


@dataclass(frozen=True)
class CompetencyQuestionBank:
    questions: tuple[CompetencyQuestion, ...]

    def by_number(self, number: int) -> CompetencyQuestion | None:
        return next((question for question in self.questions if question.number == number), None)


@lru_cache(maxsize=4)
def _load_cached(path: str) -> CompetencyQuestionBank:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        headers = tuple(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
        if headers != EXPECTED_HEADERS:
            raise ValueError(f"질문 은행 컬럼이 예상과 다릅니다: {headers}")

        questions = []
        seen_numbers: set[int] = set()
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            number = int(row[0])
            if number in seen_numbers:
                raise ValueError(f"질문 번호가 중복되었습니다: {number}")
            seen_numbers.add(number)
            questions.append(CompetencyQuestion(
                number=number,
                category=str(row[1] or "").strip(),
                competency=str(row[2] or "").strip(),
                competency_description=str(row[3] or "").strip(),
                question_type=str(row[4] or "").strip(),
                text=str(row[5] or "").strip(),
                follow_ups=_lines(row[6]),
                intents=_lines(row[7]),
                positive_checkpoints=_lines(row[8]),
                negative_checkpoints=_lines(row[9]),
            ))

        if not questions:
            raise ValueError("질문 은행이 비어 있습니다.")
        return CompetencyQuestionBank(tuple(questions))
    finally:
        workbook.close()


def load_competency_question_bank(path: str | Path = WORKBOOK_PATH) -> CompetencyQuestionBank:
    return _load_cached(str(Path(path).resolve()))